#!/usr/bin/env python3
"""
Build app/static/data/biodiversity.json from the Biological Profile workbook.

    python scripts/build_biodiversity.py Biological_Profile_Report_data.xlsx

This replaces the browser-side parse that used to read the whole 380 KB .xlsx
from /public/data on every visit. The output JSON is about 60 KB, keeps the raw
report off the public web root, and matches how every other dataset on the site
is served.

Three sheet layouts contribute, and sheets matching none of them are skipped, so
site-percentage tables, density summaries and image placeholders never appear:

  A. COLUMN BLOCKS - algae, zooplankton, macrobenthos. Column A holds "Class",
     the row beneath it holds one class per column, then "Species" and that
     column's species below. Detected automatically. Reading stops at the first
     fully blank row, because each sheet ends with a caption line that would
     otherwise be read as a species.

  B. ROW TABLES - fish, birds, reptiles, amphibians. One record per row. The
     column positions differ per sheet and there is no shared schema, so
     ROW_SHEETS declares where the rank and species columns are. Only the
     positions are declared; every name still comes from the workbook.

  C. LIFE-FORM SHEETS - Lower-zone macrophytes. These have no rank column at
     all: the sheet itself is the life form.

Plus D, composition-only sheets: a group recorded for a zone as proportions with
no species names. These cannot be tree leaves so they are returned separately.

Output:

    {
      "zones": [ { "zoneKey": "upper",
                   "groups": [ { "name", "label", "rank",
                                 "classes": [ { "name", "species": [...] } ] } ] } ],
      "composition": { "lower": [ { "group",
                                    "rows": [ {"name","value","share","notes"} ] } ] }
    }
"""

import json
import re
import sys
from pathlib import Path

from openpyxl import load_workbook

ZONE_RE = re.compile(r"\b(upper|middle|lower)\b", re.I)
ZONE_ORDER = ["upper", "middle", "lower"]
ZONE_TAG = re.compile(r"\((UZ|MZ|LZ)\)", re.I)
ZONE_TAG_MAP = {"UZ": "upper", "MZ": "middle", "LZ": "lower"}

# A rank cell longer than this is the sheet's trailing caption, not a taxon.
MAX_RANK_LEN = 40

# sheet key (trimmed, lower-cased) -> group, zones, rank, [(rankCol, speciesCol)]
# zones = None means the zone is tagged per row as "(UZ)"/"(MZ)" in the rank cell.
ROW_SHEETS = [
    ("details of wb birds of the uppe", "Birds", ["upper"], "Order", [(0, 3)]),
    ("wd birds upper", "Birds", ["upper"], "Order", [(0, 3)]),
    ("wb middle", "Birds", ["middle"], "Order", [(0, 3)]),
    ("details of wd birds of the midd", "Birds", ["middle"], "Order", [(0, 3)]),
    ("wa birds upper , middle", "Birds", None, "Order", [(0, 3)]),
    ("amphibians upper", "Amphibians", ["upper"], "Order", [(2, 1)]),
    ("amphibians middle", "Amphibians", ["middle"], "Order", [(2, 1)]),
    ("reptiles upper", "Reptiles", ["upper"], "Order", [(2, 1)]),
    ("reptiles middle", "Reptiles", ["middle"], "Order", [(2, 1)]),
    ("repltiles lower", "Reptiles", ["lower"], "Order", [(2, 1)]),
    # Fish are keyed on Family throughout: the "Common fishes" sheet carries
    # both Family and Order, and Family is what the other two fish sheets use.
    ("common fishes in the upper and", "Fish", ["upper", "middle"], "Family", [(0, 2), (3, 5)]),
    ("consolidated fish diversity", "Fish", ["lower"], "Family", [(1, 2)]),
    ("fish in narmada", "Fish", ["upper", "middle", "lower"], "Family", [(0, 1)]),
]

# sheet key -> life form. A row counts only when the density column is filled,
# which drops each sheet's caption line.
LIFEFORM_SHEETS = [
    ("major tree species density and", "Trees"),
    ("dominant shrub species density", "Shrubs"),
    ("dominant herb species density", "Herbs"),
    ("major grass species density low", "Grasses"),
    ("sedge species density", "Sedges"),
    ("climbers species density lower", "Climbers"),
    ("semi aquatic lower", "Aquatic & semi-aquatic"),
]
LIFEFORM_GROUP = "Macrophytes"
LIFEFORM_ZONE = "lower"
LIFEFORM_RANK = "Life form"

# sheet key -> zone, group, (nameCol, valueCol, notesCol)
COMPOSITION_SHEETS = [
    ("macrobenthos group table", "lower", "Macrobenthos", (0, 1, 2)),
]

# Display names for groups whose sheet name is abbreviated.
GROUP_LABELS = {
    "Phyto": "Phytoplankton",
    "Semi Aquatic": "Semi-aquatic",
}


def clean(value):
    if value is None:
        return None
    text = " ".join(str(value).split())
    return text or None


def title_case(text):
    return " ".join(w[:1].upper() + w[1:] for w in text.split() if w)


def sort_key(text):
    """Approximate the browser's localeCompare: case-insensitive, then exact."""
    return (text.lower(), text)


def rows_of(worksheet):
    return [
        [clean(c) for c in row] if isinstance(row, tuple) else []
        for row in worksheet.iter_rows(values_only=True)
    ]


def find_sheet(workbook, key):
    return next((n for n in workbook.sheetnames if n.strip().lower() == key), None)


def build(path):
    workbook = load_workbook(path, data_only=True)

    # zoneKey -> groupName -> {"rank": str, "values": {rankValue: set(species)}}
    tree = {}

    def bucket(zone_key, group, rank, rank_value):
        groups = tree.setdefault(zone_key, {})
        entry = groups.setdefault(group, {"rank": rank, "values": {}})
        return entry["values"].setdefault(rank_value, set())

    # ── A. column-block sheets ───────────────────────────────────────────────
    for sheet_name in workbook.sheetnames:
        rows = rows_of(workbook[sheet_name])
        class_row = next((i for i, r in enumerate(rows) if r and r[0] == "Class"), None)
        if class_row is None:
            continue
        species_row = next(
            (i for i, r in enumerate(rows) if i > class_row and r and r[0] == "Species"), None
        )
        if species_row is None:
            continue

        zone_match = ZONE_RE.search(sheet_name)
        if not zone_match:
            continue
        zone_key = zone_match.group(1).lower()
        group = title_case(ZONE_RE.sub(" ", sheet_name).strip())
        if not group:
            continue

        header = rows[class_row + 1] if class_row + 1 < len(rows) else []
        for col, class_name in enumerate(header):
            if not class_name:
                continue
            target = bucket(zone_key, group, "Class", class_name)
            for i in range(species_row + 1, len(rows)):
                if all(c is None for c in rows[i]):
                    break  # caption line follows
                if col < len(rows[i]) and rows[i][col]:
                    target.add(rows[i][col])

    # ── B. row tables ────────────────────────────────────────────────────────
    for key, group, zones, rank, blocks in ROW_SHEETS:
        sheet_name = find_sheet(workbook, key)
        if not sheet_name:
            print(f"  ! sheet not found, skipped: {key}", file=sys.stderr)
            continue
        for row in rows_of(workbook[sheet_name])[1:]:
            for rank_col, species_col in blocks:
                rank_value = row[rank_col] if rank_col < len(row) else None
                species = row[species_col] if species_col < len(row) else None
                if not rank_value or not species:
                    continue
                if len(rank_value) > MAX_RANK_LEN:
                    continue
                row_zones = zones
                if row_zones is None:
                    tag = ZONE_TAG.search(rank_value)
                    if not tag:
                        continue
                    row_zones = [ZONE_TAG_MAP[tag.group(1).upper()]]
                    rank_value = ZONE_TAG.sub("", rank_value).strip()
                    if not rank_value:
                        continue
                for zone_key in row_zones:
                    bucket(zone_key, group, rank, rank_value).add(species)

    # ── C. life-form sheets ──────────────────────────────────────────────────
    for key, life_form in LIFEFORM_SHEETS:
        sheet_name = find_sheet(workbook, key)
        if not sheet_name:
            print(f"  ! sheet not found, skipped: {key}", file=sys.stderr)
            continue
        for row in rows_of(workbook[sheet_name])[1:]:
            species = row[0] if row else None
            density = row[1] if len(row) > 1 else None
            if not species or not density:
                continue
            bucket(LIFEFORM_ZONE, LIFEFORM_GROUP, LIFEFORM_RANK, life_form).add(species)

    # ── D. composition-only sheets ───────────────────────────────────────────
    composition = {}
    for key, zone_key, group, (name_col, value_col, notes_col) in COMPOSITION_SHEETS:
        sheet_name = find_sheet(workbook, key)
        if not sheet_name:
            continue
        parsed = []
        for row in rows_of(workbook[sheet_name])[1:]:
            name = row[name_col] if name_col < len(row) else None
            raw = row[value_col] if value_col < len(row) else None
            if not name or raw is None or len(name) > MAX_RANK_LEN:
                continue
            try:
                value = float(raw)
            except (TypeError, ValueError):
                continue
            parsed.append(
                {
                    "name": name,
                    "value": value,
                    "notes": row[notes_col] if notes_col < len(row) else None,
                }
            )
        if not parsed:
            continue
        # The column is headed "(%)" but the values sum to 1.000, so the share is
        # derived from the data rather than reading the raw number as a percent.
        total = sum(r["value"] for r in parsed) or 1
        for r in parsed:
            r["share"] = round(r["value"] / total * 100, 4)
        composition.setdefault(zone_key, []).append({"group": group, "rows": parsed})

    # ── assemble ─────────────────────────────────────────────────────────────
    zones = []
    for zone_key, groups in tree.items():
        group_list = []
        for name, entry in groups.items():
            classes = [
                {"name": rank_value, "species": sorted(species, key=sort_key)}
                for rank_value, species in entry["values"].items()
                if species
            ]
            if not classes:
                continue
            classes.sort(key=lambda c: sort_key(c["name"]))
            group_list.append(
                {
                    "name": name,
                    "label": GROUP_LABELS.get(name, name),
                    "rank": entry["rank"],
                    "classes": classes,
                }
            )
        if not group_list:
            continue
        group_list.sort(key=lambda g: sort_key(g["name"]))
        zones.append({"zoneKey": zone_key, "groups": group_list})

    zones.sort(
        key=lambda z: ZONE_ORDER.index(z["zoneKey"]) if z["zoneKey"] in ZONE_ORDER else 99
    )
    return {"zones": zones, "composition": composition}


def main():
    if len(sys.argv) != 2:
        sys.exit("usage: build_biodiversity.py <Biological_Profile_Report_data.xlsx>")

    payload = build(sys.argv[1])

    destination = Path("app/static/data/biodiversity.json")
    destination.parent.mkdir(parents=True, exist_ok=True)
    with destination.open("w", encoding="utf-8") as handle:
        json.dump(payload, handle, ensure_ascii=False, separators=(",", ":"))

    total = 0
    for zone in payload["zones"]:
        print(f"{zone['zoneKey']}:")
        for group in zone["groups"]:
            count = sum(len(c["species"]) for c in group["classes"])
            total += count
            print(
                f"   {group['label']:<16} by {group['rank']:<10}"
                f" {len(group['classes']):>2} nodes, {count:>3} species"
            )
    for zone_key, blocks in payload["composition"].items():
        for block in blocks:
            print(f"{zone_key}: {block['group']} composition, {len(block['rows'])} groups")
    print(f"\ntotal species entries: {total}")
    print(f"wrote {destination} ({destination.stat().st_size / 1024:.1f} KB)")


if __name__ == "__main__":
    main()
